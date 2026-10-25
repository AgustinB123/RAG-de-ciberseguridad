import os
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def _agregar_texto_formateado(parrafo, texto):
    partes = re.split(r"(\*\*.*?\*\*)", texto)
    for parte in partes:
        if parte.startswith("**") and parte.endswith("**"):
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        else:
            parrafo.add_run(parte)


def generar_word(texto_markdown, nombre_archivo="outputs/Informe_Ciberseguridad.docx"):
    documento = docx.Document()
    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Calibri"
    estilo_normal.font.size = Pt(11)

    texto_limpio = re.sub(r"\\+([#*])", r"\1", texto_markdown)
    titulo_informe = "Informe de Diagnóstico de Ciberseguridad"
    titulo = documento.add_heading(titulo_informe, 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    for linea in texto_limpio.splitlines():
        linea = linea.strip()
        if not linea:
            continue

        if linea.startswith("# "):
            titulo_markdown = linea[2:].strip()
            if titulo_markdown.replace("🛡️ ", "") == titulo_informe:
                continue
            encabezado = documento.add_heading(titulo_markdown, level=1)
            encabezado.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif linea.startswith("## "):
            encabezado = documento.add_heading(linea[3:].strip(), level=2)
            encabezado.style.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        elif linea.startswith(("- ", "* ")):
            parrafo = documento.add_paragraph(style="List Bullet")
            _agregar_texto_formateado(parrafo, linea[2:])
        else:
            parrafo = documento.add_paragraph()
            _agregar_texto_formateado(parrafo, linea)

    carpeta = os.path.dirname(nombre_archivo)
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)
    documento.save(nombre_archivo)
    return nombre_archivo